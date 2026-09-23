"""Bounded document extraction. Runs in an isolated, timed subprocess."""
import json
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

SUPPORTED = {'.pdf':'application/pdf', '.docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
 '.doc':'application/msword', '.xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
 '.xls':'application/vnd.ms-excel', '.jpg':'image/jpeg', '.jpeg':'image/jpeg', '.png':'image/png'}


def redact(text):
    # Conservative removal of card-like digit sequences and IBANs before persistence/model calls.
    text = re.sub(r'\b(?:\d[ -]?){13,19}\b', '[REDACTED_PAYMENT_DATA]', text)
    return re.sub(r'\b[A-Z]{2}\d{2}[A-Z0-9]{11,30}\b', '[REDACTED_PAYMENT_DATA]', text, flags=re.I)


def validate_file(path, extension):
    head = path.read_bytes()[:16]
    signatures = {'.pdf':b'%PDF-', '.doc':bytes.fromhex('D0CF11E0A1B11AE1'), '.xls':bytes.fromhex('D0CF11E0A1B11AE1'),
                  '.jpg':b'\xff\xd8\xff', '.jpeg':b'\xff\xd8\xff', '.png':b'\x89PNG\r\n\x1a\n'}
    if extension in signatures and not head.startswith(signatures[extension]):
        raise ValueError('File signature does not match extension')
    if extension in ('.docx','.xlsx'):
        with zipfile.ZipFile(path) as z:
            info=z.infolist()
            if len(info)>3000 or sum(x.file_size for x in info)>50*1024*1024:
                raise ValueError('Expanded archive exceeds limits')
            required='word/document.xml' if extension=='.docx' else 'xl/workbook.xml'
            if required not in z.namelist() or any('vbaproject' in x.filename.lower() for x in info):
                raise ValueError('Unsupported Office document structure or macros')


def run(command):
    return subprocess.run(command, check=True, capture_output=True, text=True, timeout=30).stdout


def ocr(path, languages):
    return run(['tesseract',str(path),'stdout','-l',languages])


def extract(path, languages):
    ext=path.suffix.lower(); validate_file(path, ext)
    blocks=[]; warnings=[]
    def add(text, **source):
        if str(text).strip():
            if len(str(text))>4000:
                warnings.append('Фрагмент текста сокращён до 4000 символов; для полного анализа разделите документ.')
            blocks.append({'text':redact(str(text))[:4000], 'source':source})
    if ext in ('.jpg','.jpeg','.png'):
        from PIL import Image
        Image.MAX_IMAGE_PIXELS=20_000_000
        with Image.open(path) as image:
            if image.width*image.height>20_000_000:
                raise ValueError('Image exceeds 20 megapixels')
            image.verify()
        add(ocr(path,languages), page=1)
        warnings.append('Фото обработано OCR. Визуальная идентификация без читаемой маркировки требует уточнения.')
    elif ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(path)
        if reader.is_encrypted:
            raise ValueError('Password-protected PDF is unsupported')
        if len(reader.pages)>20:
            raise ValueError('PDF limit is 20 pages; split the document')
        for index,page in enumerate(reader.pages):
            text=page.extract_text() or ''
            if len(text.strip())<20:
                with tempfile.TemporaryDirectory() as d:
                    target=Path(d)/'page'
                    run(['pdftoppm','-f',str(index+1),'-l',str(index+1),'-scale-to','2000','-singlefile','-png',str(path),str(target)])
                    text=ocr(target.with_suffix('.png'),languages)
            add(text,page=index+1)
    elif ext=='.docx':
        from docx import Document
        doc=Document(path)
        for i,p in enumerate(doc.paragraphs): add(p.text,paragraph=i+1)
        for table_no,table in enumerate(doc.tables):
            for row_no,row in enumerate(table.rows): add(' | '.join(c.text for c in row.cells),table=table_no+1,row=row_no+1)
    elif ext=='.doc':
        with tempfile.TemporaryDirectory() as d:
            run(['libreoffice','-env:UserInstallation=file://'+d+'/profile','--headless','--convert-to','docx','--outdir',d,str(path)])
            return extract(Path(d)/(path.stem+'.docx'),languages)
    elif ext=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=True,keep_links=False)
        try:
            for sheet in book:
                for i,row in enumerate(sheet.iter_rows(values_only=True)):
                    if i>=500 or len(blocks)>=500: raise ValueError('Spreadsheet exceeds 500 rows; split the document')
                    add(' | '.join('' if v is None else str(v) for v in row),sheet=sheet.title,row=i+1)
        finally: book.close()
        warnings.append('Использованы сохранённые значения формул; формулы не выполнялись.')
    elif ext=='.xls':
        import xlrd
        book=xlrd.open_workbook(path,on_demand=True)
        try:
            for sheet in book.sheets():
                if sheet.nrows>500 or len(blocks)+sheet.nrows>500: raise ValueError('Spreadsheet exceeds 500 rows; split the document')
                for i in range(sheet.nrows): add(' | '.join(str(v) for v in sheet.row_values(i)),sheet=sheet.name,row=i+1)
        finally: book.release_resources()
    else: raise ValueError('Unsupported file type')
    if len(blocks)>500 or sum(len(b['text']) for b in blocks)>40000:
        raise ValueError('Extracted content exceeds limit; split the document')
    if not blocks: warnings.append('Не удалось прочитать текст. Уточните артикул или загрузите более чёткое фото.')
    return {'blocks':blocks,'warnings':warnings}


if __name__=='__main__':
    try:
        print(json.dumps(extract(Path(sys.argv[1]),sys.argv[2]),ensure_ascii=False))
    except Exception as exc:
        # Do not log document contents or library exception payloads.
        print(json.dumps({'error':'Unable to extract document: unsupported, unreadable, encrypted, or over processing limits','error_type':type(exc).__name__}))
        sys.exit(1)
