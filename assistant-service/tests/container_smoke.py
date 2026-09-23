"""Real HTTP + worker + OCR/Office smoke check. Run inside the built image."""
import os
from pathlib import Path
import subprocess
import tempfile
import time

import httpx
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont

with tempfile.TemporaryDirectory() as directory:
    folder=Path(directory)
    env={**os.environ,'ASSISTANT_DATA_DIR':str(folder/'data'),'ASSISTANT_ALLOW_DEMO_SESSIONS':'true',
         'ASSISTANT_CATALOG_MODE':'demo','ASSISTANT_CART_MODE':'demo','ASSISTANT_MODEL_API_KEY':'','ASSISTANT_MODEL_NAME':''}
    api=subprocess.Popen(['uvicorn','app.main:app','--host','127.0.0.1','--port','8123'],env=env)
    worker=subprocess.Popen(['python','-m','app.worker'],env=env)
    try:
        with httpx.Client(base_url='http://127.0.0.1:8123',timeout=10) as client:
            for _ in range(100):
                try:
                    if client.get('/health/live').status_code==200: break
                except httpx.ConnectError: pass
                time.sleep(.1)
            else: raise AssertionError('API did not start')
            data=client.post('/v1/sessions',json={}).json()
            base='/v1/sessions/'+data['session_id']; headers={'Authorization':'Bearer '+data['session_token']}
            doc=Document(); doc.add_paragraph('DEMO-C16 quantity 2'); doc.save(folder/'spec.docx')
            wb=Workbook(); wb.active.append(['DEMO-C16',2]); wb.save(folder/'spec.xlsx')
            image=Image.new('RGB',(1200,220),'white'); ImageDraw.Draw(image).text((30,60),'DEMO-C16 quantity 2',font=ImageFont.load_default(size=54),fill='black')
            image.save(folder/'photo.jpg'); image.save(folder/'photo.png'); image.save(folder/'scan.pdf')
            for kind in ('doc','pdf'):
                subprocess.run(['libreoffice','-env:UserInstallation=file://'+directory+'/lo-profile','--headless','--convert-to',kind,'--outdir',directory,str(folder/'spec.docx')],check=True,capture_output=True,timeout=30)
            inputs=[folder/'spec.docx',folder/'spec.doc',folder/'spec.xlsx',Path('/checks/fixtures/spec.xls'),folder/'spec.pdf',folder/'scan.pdf',folder/'photo.jpg',folder/'photo.png']
            attachment_ids=[]
            for path in inputs:
                response=client.post(base+'/attachments',headers=headers,files={'file':(path.name,path.read_bytes())})
                assert response.status_code==202,response.text
                attachment_ids.append((path.name,response.json()['attachment_id']))
            deadline=time.monotonic()+120
            remaining=dict(attachment_ids)
            while remaining and time.monotonic()<deadline:
                for name,aid in list(remaining.items()):
                    state=client.get(base+'/attachments/'+aid,headers=headers).json()
                    assert state['status']!='failed',(name,state)
                    if state['status']=='ready':
                        assert state['blocks_count']>0,(name,state)
                        print('EXTRACTION OK',name,flush=True)
                        del remaining[name]
                time.sleep(.2)
            assert not remaining,remaining
            start=time.monotonic()
            response=client.post(base+'/messages',headers=headers,json={'client_message_id':'smoke','text':'DEMO-C16','attachment_ids':[attachment_ids[0][1]]})
            state=response.json()
            while state['status'] not in ('completed','failed') and time.monotonic()-start<15:
                time.sleep(.1); state=client.get(state['status_url'],headers=headers).json()
            assert state['status']=='completed' and state['products'],state
            print('HTTP message complete in',round(time.monotonic()-start,3),'seconds',flush=True)
            proposal=client.post(base+'/cart/proposals',headers=headers,json={'items':[{'product_id':'demo-1','warehouse_id':'demo-almaty','quantity':'2'}]}).json()
            result=client.post(base+'/cart/proposals/'+proposal['proposal_id']+'/confirm',headers={**headers,'Idempotency-Key':'smoke'},json={'proposal_version':1,'confirmed':True,'confirmation_token':proposal['confirmation_token']})
            assert result.status_code==200,result.text
            assert result.json()['cart']['mode']=='demo'
            print('PASS: real API + persistent worker + all 8 document variants + confirmed demo cart',flush=True)
    finally:
        worker.terminate(); api.terminate()
        worker.wait(timeout=10); api.wait(timeout=10)
