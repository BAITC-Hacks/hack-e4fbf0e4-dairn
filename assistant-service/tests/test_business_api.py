import asyncio
import io
import json
import time
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook
from docx import Document

from app.core.config import Settings
from app.main import create_app
from app.worker import process_one


@pytest.fixture
async def api(tmp_path):
    settings=Settings(environment='test',data_dir=tmp_path,message_wait_seconds=0)
    app=create_app(settings)
    async with app.router.lifespan_context(app):
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://test') as client:
            yield client,app.state.runtime


async def session(client):
    response=await client.post('/v1/sessions',json={})
    assert response.status_code==201,response.text
    data=response.json()
    return '/v1/sessions/'+data['session_id'],{'Authorization':'Bearer '+data['session_token']}


async def drain(rt,limit=20):
    for _ in range(limit):
        if not await process_one(rt): break


@pytest.mark.asyncio
async def test_text_context_idempotency_and_sse(api):
    client,rt=api; base,headers=await session(client)
    payload={'client_message_id':'one','text':'DEMO-C16-OLD'}
    first=await client.post(base+'/messages',json=payload,headers=headers)
    assert first.status_code==202
    await drain(rt)
    done=(await client.get(first.json()['status_url'],headers=headers)).json()
    assert done['status']=='completed' and done['alternatives']
    assert 'Демонстрацион' in done['answer']
    replay=await client.post(base+'/messages',json=payload,headers=headers)
    assert replay.status_code==200 and replay.json()['message_id']==first.json()['message_id']
    conflict=await client.post(base+'/messages',json={**payload,'text':'changed'},headers=headers)
    assert conflict.status_code==409
    events=await client.get(done['events_url'],headers=headers)
    assert 'event: completed' in events.text
    follow=await client.post(base+'/messages',json={'client_message_id':'two','text':'А характеристики?'},headers=headers)
    await drain(rt)
    result=(await client.get(follow.json()['status_url'],headers=headers)).json()
    assert result['products']
    assert (await client.get(base+'/cart',headers=headers)).json()['items']==[]


@pytest.mark.asyncio
async def test_session_isolation_and_expiry(api):
    client,rt=api; base,headers=await session(client); other,other_headers=await session(client)
    msg=(await client.post(base+'/messages',headers=headers,json={'client_message_id':'a','text':'hello'})).json()
    assert (await client.get(msg['status_url'],headers=other_headers)).status_code==403
    assert (await client.get(other+'/messages/'+msg['message_id'],headers=other_headers)).status_code==404
    assert (await client.get(base+'/cart')).status_code==401
    with rt.store.transaction() as con: con.execute("UPDATE objects SET expires=0 WHERE kind='session' AND session=?",(base.split('/')[-1],))
    assert (await client.get(base+'/cart',headers=headers)).status_code==401


@pytest.mark.asyncio
@pytest.mark.parametrize('kind',['xlsx','docx'])
async def test_documents_and_attachment_only_message(api,kind):
    client,rt=api; base,headers=await session(client); buffer=io.BytesIO()
    if kind=='xlsx':
        wb=Workbook(); wb.active.append(['Артикул','Количество']); wb.active.append(['DEMO-C16',2]); wb.save(buffer)
    else:
        doc=Document(); doc.add_paragraph('DEMO-C16'); doc.save(buffer)
    uploaded=await client.post(base+'/attachments',headers=headers,files={'file':('spec.'+kind,buffer.getvalue(),'application/octet-stream')})
    assert uploaded.status_code==202,uploaded.text
    attachment=uploaded.json()
    message=await client.post(base+'/messages',headers=headers,json={'client_message_id':'file','attachment_ids':[attachment['attachment_id']]})
    await drain(rt)
    ready=(await client.get(base+'/attachments/'+attachment['attachment_id'],headers=headers)).json()
    assert ready['status']=='ready',ready
    assert ready['blocks_count']>0
    assert not list((rt.settings.data_dir/'uploads').iterdir())
    result=(await client.get(message.json()['status_url'],headers=headers)).json()
    assert result['status']=='completed' and result['products'],result


@pytest.mark.asyncio
async def test_bad_files_fail_without_leaking_content(api):
    client,rt=api; base,headers=await session(client)
    response=await client.post(base+'/attachments',headers=headers,files={'file':('attack.exe',b'x')})
    assert response.status_code==415
    response=await client.post(base+'/attachments',headers=headers,files={'file':('fake.pdf',b'PRIVATE CONTENT')})
    await drain(rt)
    data=(await client.get(base+'/attachments/'+response.json()['attachment_id'],headers=headers)).json()
    assert data['status']=='failed'
    assert 'PRIVATE CONTENT' not in json.dumps(data)
    assert not list((rt.settings.data_dir/'uploads').iterdir())


@pytest.mark.asyncio
async def test_cart_explicit_confirmation_and_replay(api):
    client,rt=api; base,headers=await session(client)
    selection={'items':[{'product_id':'demo-1','warehouse_id':'demo-almaty','quantity':'2'}]}
    proposal=(await client.post(base+'/cart/proposals',headers=headers,json=selection)).json()
    assert (await client.get(base+'/cart',headers=headers)).json()['items']==[]
    body={'proposal_version':proposal['version'],'confirmation_token':proposal['confirmation_token'],'confirmed':True}
    url=base+'/cart/proposals/'+proposal['proposal_id']+'/confirm'
    assert (await client.post(url,headers={**headers,'Idempotency-Key':'x'},json={**body,'confirmed':False})).status_code==422
    result=await client.post(url,headers={**headers,'Idempotency-Key':'x'},json=body)
    assert result.status_code==200,result.text
    replay=await client.post(url,headers={**headers,'Idempotency-Key':'x'},json=body)
    assert replay.json()==result.json()
    assert result.json()['cart']['items'][0]['selection']['quantity']=='2'
    assert 'session_id=' in result.json()['cart']['cart_url']
    assert (await rt.catalog.detail('demo-1'))['stock'][0]['available_quantity']=='18'
    assert (await client.post(url,headers={**headers,'Idempotency-Key':'another'},json=body)).status_code==409


@pytest.mark.asyncio
async def test_concurrent_confirmations_cannot_oversell(api):
    client,rt=api
    async def prepare():
        base,headers=await session(client)
        proposal=(await client.post(base+'/cart/proposals',headers=headers,json={'items':[{'product_id':'demo-1','warehouse_id':'demo-almaty','quantity':'15'}]})).json()
        return base,headers,proposal
    prepared=[await prepare(),await prepare()]
    async def commit(base,headers,p):
        return await client.post(base+'/cart/proposals/'+p['proposal_id']+'/confirm',headers={**headers,'Idempotency-Key':'commit'},json={'proposal_version':1,'confirmation_token':p['confirmation_token'],'confirmed':True})
    results=await asyncio.gather(*(commit(*p) for p in prepared))
    assert sorted(r.status_code for r in results)==[200,409]
    assert (await rt.catalog.detail('demo-1'))['stock'][0]['available_quantity']=='5'


@pytest.mark.asyncio
async def test_superseded_and_wrong_confirmation(api):
    client,rt=api; base,headers=await session(client)
    selection={'items':[{'product_id':'demo-1','warehouse_id':'demo-almaty','quantity':'1'}]}
    first=(await client.post(base+'/cart/proposals',headers=headers,json=selection)).json()
    second=(await client.post(base+'/cart/proposals',headers=headers,json=selection)).json()
    def body(p):return {'proposal_version':1,'confirmation_token':p['confirmation_token'],'confirmed':True}
    assert (await client.post(base+'/cart/proposals/'+first['proposal_id']+'/confirm',headers={**headers,'Idempotency-Key':'old'},json=body(first))).status_code==409
    assert (await client.post(base+'/cart/proposals/'+second['proposal_id']+'/confirm',headers={**headers,'Idempotency-Key':'bad'},json={**body(second),'confirmation_token':'wrong'})).status_code==403


@pytest.mark.asyncio
async def test_model_request_single_call_and_timeout_fallback(api):
    client,rt=api; base,headers=await session(client)
    from pydantic import SecretStr
    rt.settings.model_api_key=SecretStr('test-only'); rt.settings.model_name='configured-test-model'
    calls=[]
    async def respond(request):
        body=json.loads(request.content); calls.append(body)
        assert body['store'] is False and 'tools' not in body
        raise httpx.ReadTimeout('mock timeout')
    await rt.model.client.aclose()
    rt.model.client=httpx.AsyncClient(transport=httpx.MockTransport(respond))
    msg=(await client.post(base+'/messages',headers=headers,json={'client_message_id':'m','text':'Подберите автомат C16'})).json()
    await drain(rt)
    result=(await client.get(msg['status_url'],headers=headers)).json()
    assert result['status']=='completed' and result['answer_mode']=='grounded_template'
    assert len(calls)==1 and any('Модель' in w for w in result['warnings'])


@pytest.mark.asyncio
async def test_durable_queue_recovery_after_restart(api):
    client,rt=api; base,headers=await session(client)
    msg=(await client.post(base+'/messages',headers=headers,json={'client_message_id':'persist','text':'DEMO-C16'})).json()
    job=rt.store.claim()
    with rt.store.transaction() as con: con.execute('UPDATE jobs SET lease=0 WHERE id=?',(job['id'],))
    from app.runtime import Runtime
    restarted=Runtime(rt.settings)
    try: await drain(restarted)
    finally: await restarted.close()
    assert (await client.get(msg['status_url'],headers=headers)).json()['status']=='completed'


@pytest.mark.asyncio
async def test_bootstrap_required_and_openapi(tmp_path):
    from pydantic import SecretStr
    from openapi_spec_validator import validate
    settings=Settings(environment='production',data_dir=tmp_path,allow_demo_sessions=False,bootstrap_token=SecretStr('server-only'))
    app=create_app(settings); validate(app.openapi())
    async with app.router.lifespan_context(app):
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app),base_url='http://test') as client:
            assert (await client.post('/v1/sessions',json={})).status_code==401
            assert (await client.post('/v1/sessions',json={},headers={'Authorization':'Bearer server-only'})).status_code==201


@pytest.mark.asyncio
async def test_attachment_ownership_pending_and_failure(api):
    client,rt=api; base,headers=await session(client); other,other_headers=await session(client)
    file=(await client.post(base+'/attachments',headers=headers,files={'file':('broken.pdf',b'broken')})).json()
    denied=await client.post(other+'/messages',headers=other_headers,json={'client_message_id':'stolen','attachment_ids':[file['attachment_id']]})
    assert denied.status_code==404
    msg=(await client.post(base+'/messages',headers=headers,json={'client_message_id':'waiting','attachment_ids':[file['attachment_id']]})).json()
    # Force message to be claimed before the upload, proving that pending files are accepted.
    with rt.store.transaction() as con:
        con.execute('UPDATE jobs SET available=0 WHERE id=?',(msg['message_id'],))
    await process_one(rt)
    pending=(await client.get(msg['status_url'],headers=headers)).json()
    assert pending['status']=='waiting_for_attachments'
    await drain(rt)
    await asyncio.sleep(0.35)
    await drain(rt)
    assert (await client.get(msg['status_url'],headers=headers)).json()['error']['code']=='attachment_failed'


@pytest.mark.asyncio
async def test_cleanup_and_card_redaction(api):
    client,rt=api; base,headers=await session(client)
    msg=(await client.post(base+'/messages',headers=headers,json={'client_message_id':'private','text':'4111 1111 1111 1111'})).json()
    assert '4111' not in msg['text']
    file=(await client.post(base+'/attachments',headers=headers,files={'file':('test.pdf',b'%PDF-x')})).json()
    with rt.store.transaction() as con:
        con.execute('UPDATE objects SET expires=0 WHERE session=?',(base.split('/')[-1],))
    rt.store.cleanup()
    assert not list((rt.settings.data_dir/'uploads').iterdir())
    assert (await client.get(msg['status_url'],headers=headers)).status_code==401


@pytest.mark.asyncio
async def test_size_limits_and_cart_strict_boolean(api):
    client,rt=api; base,headers=await session(client)
    rt.settings.upload_max_bytes=10
    result=await client.post(base+'/attachments',headers=headers,files={'file':('big.pdf',b'%PDF-'+b'x'*20)})
    assert result.status_code==413
    result=await client.post(base+'/cart/proposals/x/confirm',headers={**headers,'Idempotency-Key':'x'},json={'confirmed':1,'proposal_version':1,'confirmation_token':'x'})
    assert result.status_code==422


@pytest.mark.asyncio
async def test_exact_sku_bypasses_model(api):
    client,rt=api; base,headers=await session(client)
    async def forbidden(evidence):
        raise AssertionError('Exact SKU lookup should not invoke a model')
    rt.model.answer=forbidden
    msg=(await client.post(base+'/messages',headers=headers,json={'client_message_id':'fast','text':'DEMO-C16'})).json()
    await drain(rt)
    assert (await client.get(msg['status_url'],headers=headers)).json()['status']=='completed'


@pytest.mark.asyncio
async def test_remote_catalog_failure_is_explicit(api):
    client,rt=api; base,headers=await session(client)
    rt.settings.catalog_mode='http'; rt.settings.cart_mode='disabled'
    async def fail(request): return httpx.Response(503,json={})
    await rt.catalog.client.aclose()
    rt.catalog.client=httpx.AsyncClient(base_url='http://catalog/',transport=httpx.MockTransport(fail))
    msg=(await client.post(base+'/messages',headers=headers,json={'client_message_id':'remote','text':'DEMO-C16'})).json()
    await drain(rt)
    result=(await client.get(msg['status_url'],headers=headers)).json()
    assert result['status']=='completed' and not result['products']
    assert any('недоступен' in w for w in result['warnings'])
    assert (await client.get(base+'/cart',headers=headers)).status_code==503
